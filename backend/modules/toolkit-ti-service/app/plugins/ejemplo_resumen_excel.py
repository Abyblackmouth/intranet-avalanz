"""Plugin de ejemplo: resume las hojas de un Excel."""
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

FICHA = {
    "nombre": "Resumen de un Excel",
    "descripcion": "Recibe un Excel y entrega otro con sus hojas, filas, columnas y encabezados. Sirve para probar el motor de plugins.",
    "tipo": "manual",
    "grupo": "Ejemplos",
    "parametros": [
        {"nombre": "archivo", "tipo": "archivo", "extensiones": [".xlsx"], "requerido": True, "ayuda": "Cualquier archivo .xlsx"},
        {"nombre": "incluir_encabezados", "tipo": "booleano", "ayuda": "Agregar la primera fila de cada hoja"},
    ],
    "tiempo_maximo": 60,
}


def ejecutar(params: dict, archivos: dict, salida: Path, log) -> None:
    origen = archivos["archivo"]
    log(f"Leyendo {origen.name}")
    libro = load_workbook(origen, read_only=True, data_only=True)
    res = Workbook()
    hoja = res.active
    hoja.title = "Resumen"
    columnas = ["Hoja", "Filas", "Columnas"] + (["Encabezados"] if params.get("incluir_encabezados") else [])
    hoja.append(columnas)
    for c in hoja[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1A4FA0")
    for ws in libro.worksheets:
        fila = [ws.title, ws.max_row or 0, ws.max_column or 0]
        if params.get("incluir_encabezados"):
            primera = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
            fila.append(", ".join(str(v) for v in primera if v is not None))
        hoja.append(fila)
        log(f"  {ws.title}: {fila[1]} filas, {fila[2]} columnas")
    for col, ancho in zip("ABCD", (28, 10, 10, 60)):
        hoja.column_dimensions[col].width = ancho
    destino = salida / f"resumen_{origen.stem}.xlsx"
    res.save(destino)
    log(f"Generado {destino.name}")
