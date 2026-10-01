"""Alta masiva de empleados (Administración → Usuarios), solo super admin.

Layout generado al momento -> revisión sin crear nada -> confirmación que crea en
segundo plano con las mismas funciones del alta individual y del acceso a módulos,
y manda los correos de bienvenida escalonados. Las contraseñas temporales solo
viven en memoria hasta que sale su correo."""
import asyncio
import io
import logging
import re
import uuid
from datetime import datetime, timedelta, timezone

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.database import AsyncSessionFactory, get_db
from app.models.admin_models import Company, Module, ModuleRole, User
from app.routes.users import validator
from app.services import user_service

router = APIRouter(prefix="/users/bulk", tags=["users-bulk"])
log = logging.getLogger("alta_masiva")

SOLO_ADMIN = validator.require_roles(["super_admin"])
PAUSA_CORREOS = 3          # segundos entre correos de bienvenida (Office 365)
MAX_FILAS = 1000
VIGENCIA_LOTE = timedelta(hours=1)
SOLICITANTE = "Solicitante"
BASE = ["Empresa", "Nombre completo", "Correo", "Matrícula", "Puesto", "Departamento", "Teléfono"]
OBLIGATORIAS = set(BASE[:6])
PREFIJO_MODULO = "Acceso · "

LOTES: dict[str, dict] = {}       # revisiones pendientes de confirmar
TRABAJOS: dict[str, dict] = {}    # altas en proceso o terminadas

AZUL = "1A4FA0"
_fuente, _cab = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True, color="FFFFFF")
_azul, _gris = PatternFill("solid", fgColor=AZUL), PatternFill("solid", fgColor="64748B")
_linea = Side(style="thin", color="CBD5E1")
_borde = Border(left=_linea, right=_linea, top=_linea, bottom=_linea)


async def _catalogos(db: AsyncSession) -> dict:
    empresas = (await db.execute(select(Company).where(Company.is_deleted == False, Company.is_active == True)
                                 .order_by(Company.nombre_comercial))).scalars().all()
    deptos = sorted({d for (d,) in (await db.execute(select(User.departamento).where(User.is_deleted == False, User.departamento.isnot(None)))).all() if d})
    modulos = (await db.execute(select(Module).where(Module.is_deleted == False, Module.is_active == True).order_by(Module.name))).scalars().all()
    roles = (await db.execute(select(ModuleRole).where(ModuleRole.is_deleted == False, ModuleRole.is_active == True)
                              .order_by(ModuleRole.name))).scalars().all()
    por_modulo = {m.id: [r for r in roles if r.module_id == m.id] for m in modulos}
    return {"empresas": empresas, "deptos": deptos, "modulos": modulos, "roles": por_modulo}


def _xlsx(wb: Workbook, nombre: str) -> StreamingResponse:
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{nombre}"'})


def _encabezado(ws, fila: int, col: int, texto: str, obligatorio: bool, ancho: float, nota: str | None = None):
    c = ws.cell(row=fila, column=col, value=texto)
    c.font, c.fill, c.border = _cab, (_azul if obligatorio else _gris), _borde
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if nota:
        c.comment = Comment(nota, "Intranet Avalanz")
    ws.column_dimensions[c.column_letter].width = ancho


@router.get("/layout")
async def descargar_layout(db: AsyncSession = Depends(get_db), payload=Depends(SOLO_ADMIN)):
    cat = await _catalogos(db)
    wb = Workbook()
    ins = wb.active; ins.title = "Instrucciones"
    ins.sheet_view.showGridLines = False
    ins.column_dimensions["A"].width = 3
    ins["B2"] = "Alta masiva de empleados · Intranet Avalanz"; ins["B2"].font = Font(name="Arial", size=14, bold=True, color=AZUL)
    pasos = [
        "1. Llena la hoja «Empleados»: un renglón por persona, desde el renglón 2. No cambies ni muevas los encabezados.",
        "2. Las columnas azules son obligatorias. Teléfono y los accesos (gris) son opcionales.",
        "3. Empresa, Departamento y cada acceso tienen lista desplegable (hoja «Catálogos»).",
        "4. Accesos: «Solicitante» da acceso al módulo para levantar tickets; un rol da además sus permisos. Vacío = sin acceso.",
        "5. Matrícula y Teléfono son de texto: escribe los ceros a la izquierda tal cual (012185).",
        f"6. Hasta {MAX_FILAS:,} empleados por archivo. Al subirlo verás una revisión renglón por renglón antes de crear a nadie.",
    ]
    for i, t in enumerate(pasos, 4):
        ins.cell(row=i, column=2, value=t).font = _fuente
    ins.cell(row=11, column=2, value="Ejemplo de un renglón bien llenado (no lo copies a «Empleados»):").font = Font(name="Arial", size=10, bold=True)
    ejemplo = ["AGIM", "FELIPE GONZÁLEZ MARTÍNEZ", "felipe_gonzalez@avalanz.com", "012213", "ABOGADO SR", "LEGAL", "8112345678"] + [SOLICITANTE] * len(cat["modulos"])
    titulos = BASE + [f"{PREFIJO_MODULO}{m.name}" for m in cat["modulos"]]
    for i, (t, v) in enumerate(zip(titulos, ejemplo), 2):
        _encabezado(ins, 12, i, t, t in OBLIGATORIAS, 24 if i > 2 else 26)
        c = ins.cell(row=13, column=i, value=v); c.font, c.border = Font(name="Arial", size=10, italic=True, color="334155"), _borde

    ws = wb.create_sheet("Empleados")
    notas = {"Empresa": "Nombre comercial, de la lista.", "Correo": "Correo de trabajo. No puede existir ya en la intranet.",
             "Matrícula": "Con ceros a la izquierda (012185).", "Departamento": "De la lista. Si es nuevo, escríbelo: se confirmará al subir.",
             "Teléfono": "Opcional. 10 dígitos, sin espacios."}
    anchos = [24, 38, 36, 14, 34, 26, 16]
    for i, t in enumerate(BASE, 1):
        _encabezado(ws, 1, i, t + (" *" if t in OBLIGATORIAS else ""), t in OBLIGATORIAS, anchos[i - 1], notas.get(t))
    for j, m in enumerate(cat["modulos"], len(BASE) + 1):
        _encabezado(ws, 1, j, f"{PREFIJO_MODULO}{m.name}", False, 26, f"Vacío: sin acceso. «{SOLICITANTE}»: puede levantar tickets. O uno de sus roles.")
    ws.freeze_panes = "A2"; ws.row_dimensions[1].height = 32
    ultima = len(BASE) + len(cat["modulos"])
    for r in range(2, MAX_FILAS + 2):
        for col in range(1, ultima + 1):
            c = ws.cell(row=r, column=col); c.font, c.border = _fuente, _borde
        ws.cell(row=r, column=4).number_format = "@"
        ws.cell(row=r, column=7).number_format = "@"

    cs = wb.create_sheet("Catálogos")
    listas = [("Empresas", [e.nombre_comercial for e in cat["empresas"]]), ("Departamentos", cat["deptos"])]
    listas += [(f"Roles · {m.name}", [SOLICITANTE] + [r.name for r in cat["roles"][m.id]]) for m in cat["modulos"]]
    for col, (titulo, valores) in enumerate(listas, 1):
        _encabezado(cs, 1, col, titulo, True, 30)
        for i, v in enumerate(valores, 2):
            cs.cell(row=i, column=col, value=v).font = _fuente
    cs.freeze_panes = "A2"

    def lista(col_cat: int, n: int, col_emp: int, estricta: bool, titulo: str, msg: str):
        letra = cs.cell(row=1, column=col_cat).column_letter
        dv = DataValidation(type="list", formula1=f"='Catálogos'!${letra}$2:${letra}${max(n, 1) + 1}", allow_blank=True,
                            showErrorMessage=True, errorStyle="stop" if estricta else "warning", errorTitle=titulo, error=msg)
        ws.add_data_validation(dv)
        dest = ws.cell(row=1, column=col_emp).column_letter
        dv.add(f"{dest}2:{dest}{MAX_FILAS + 1}")
    lista(1, len(listas[0][1]), 1, True, "Empresa no válida", "Elige una empresa de la lista.")
    lista(2, len(listas[1][1]), 6, False, "Departamento nuevo", "No está en la lista. Si es correcto, acepta: se confirmará al subir.")
    for k, m in enumerate(cat["modulos"]):
        lista(3 + k, len(listas[2 + k][1]), len(BASE) + 1 + k, True, "Acceso no válido", f"Elige «{SOLICITANTE}» o un rol de {m.name}, o déjalo vacío.")
    wb.active = 1
    return _xlsx(wb, f"Layout_alta_masiva_{datetime.now().strftime('%Y%m%d')}.xlsx")


def _texto(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


@router.post("/revision")
async def revisar_archivo(archivo: UploadFile = File(...), db: AsyncSession = Depends(get_db), payload=Depends(SOLO_ADMIN)):
    contenido = await archivo.read()
    if len(contenido) > 5 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="El archivo pesa más de 5 MB")
    try:
        wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        ws = wb["Empleados"]
    except Exception:
        raise HTTPException(status_code=422, detail="No es el layout de alta masiva: no encontré la hoja «Empleados»")
    filas = list(ws.iter_rows(values_only=True))
    if not filas:
        raise HTTPException(status_code=422, detail="La hoja «Empleados» está vacía")
    cab = [(_texto(c).replace(" *", "")) for c in filas[0]]
    faltan = [c for c in OBLIGATORIAS if c not in cab]
    if faltan:
        raise HTTPException(status_code=422, detail=f"Al layout le faltan columnas: {', '.join(sorted(faltan))}. Descarga uno nuevo desde la intranet.")
    idx = {c: i for i, c in enumerate(cab)}

    cat = await _catalogos(db)
    empresas = {e.nombre_comercial.strip().upper(): e for e in cat["empresas"]}
    deptos = {d.upper() for d in cat["deptos"]}
    modulos = {m.name.strip().upper(): m for m in cat["modulos"]}
    cols_mod = [(i, modulos.get(c[len(PREFIJO_MODULO):].strip().upper())) for i, c in enumerate(cab) if c.startswith(PREFIJO_MODULO)]
    correos_bd = {e.lower() for (e,) in (await db.execute(select(User.email).where(User.is_deleted == False))).all()}
    matriculas_bd = {m for (m,) in (await db.execute(select(User.matricula).where(User.is_deleted == False, User.matricula.isnot(None)))).all()}

    resultado, vistos_correo, vistos_mat = [], {}, {}
    for n, fila in enumerate(filas[1:], 2):
        val = lambda c: _texto(fila[idx[c]]) if c in idx and idx[c] < len(fila) else ""
        if not any(_texto(x) for x in fila):
            continue
        if len(resultado) >= MAX_FILAS:
            raise HTTPException(status_code=422, detail=f"El archivo tiene más de {MAX_FILAS:,} empleados")
        errores, avisos, accesos = [], [], []
        empresa_txt, nombre, correo = val("Empresa").upper(), val("Nombre completo").upper(), val("Correo").lower()
        matricula, puesto, depto, telefono = val("Matrícula"), val("Puesto").upper(), val("Departamento").upper(), re.sub(r"\D", "", val("Teléfono"))
        for c, v in (("Empresa", empresa_txt), ("Nombre completo", nombre), ("Correo", correo), ("Matrícula", matricula), ("Puesto", puesto), ("Departamento", depto)):
            if not v:
                errores.append(f"falta {c.lower()}")
        empresa = empresas.get(empresa_txt)
        if empresa_txt and not empresa:
            errores.append(f"la empresa «{empresa_txt}» no existe o no está operando")
        if correo and not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", correo):
            errores.append("el correo no es válido")
        elif correo in correos_bd:
            errores.append("el correo ya existe en la intranet")
        elif correo in vistos_correo:
            errores.append(f"el correo se repite en el renglón {vistos_correo[correo]}")
        if matricula in matriculas_bd:
            errores.append("la matrícula ya existe en la intranet")
        elif matricula and matricula in vistos_mat:
            errores.append(f"la matrícula se repite en el renglón {vistos_mat[matricula]}")
        if telefono and len(telefono) != 10:
            errores.append("el teléfono debe tener 10 dígitos")
        if depto and depto not in deptos:
            avisos.append(f"departamento nuevo: «{depto}»")
        for i, mod in cols_mod:
            v = _texto(fila[i]) if i < len(fila) else ""
            if not v:
                continue
            if not mod:
                errores.append(f"la columna «{cab[i]}» no corresponde a un módulo activo")
                continue
            if v.upper() == SOLICITANTE.upper():
                accesos.append({"modulo_id": str(mod.id), "modulo": mod.name, "role_id": None, "rol": SOLICITANTE})
                continue
            rol = next((r for r in cat["roles"][mod.id] if r.name.upper() == v.upper()), None)
            if rol:
                accesos.append({"modulo_id": str(mod.id), "modulo": mod.name, "role_id": str(rol.id), "rol": rol.name})
            else:
                errores.append(f"«{v}» no es un rol de {mod.name}")
        if correo:
            vistos_correo.setdefault(correo, n)
        if matricula:
            vistos_mat.setdefault(matricula, n)
        resultado.append({"fila": n, "empresa": empresa_txt, "company_id": str(empresa.id) if empresa else None, "nombre": nombre,
                          "correo": correo, "matricula": matricula, "puesto": puesto, "departamento": depto, "telefono": telefono or None,
                          "accesos": accesos, "estado": "error" if errores else ("aviso" if avisos else "listo"),
                          "mensajes": errores + avisos})

    ahora = datetime.now(timezone.utc)
    for k in [k for k, v in LOTES.items() if v["creado"] < ahora - VIGENCIA_LOTE]:
        LOTES.pop(k, None)
    lote_id = uuid.uuid4().hex
    LOTES[lote_id] = {"creado": ahora, "filas": resultado, "por": str(payload.get("sub") or payload.get("user_id") or "")}
    cuenta = lambda e: sum(1 for f in resultado if f["estado"] == e)
    return {"lote_id": lote_id, "archivo": archivo.filename,
            "resumen": {"total": len(resultado), "listos": cuenta("listo"), "avisos": cuenta("aviso"), "errores": cuenta("error")},
            "departamentos_nuevos": sorted({f["departamento"] for f in resultado if any(m.startswith("departamento nuevo") for m in f["mensajes"])}),
            "filas": resultado}


class Confirmacion(BaseModel):
    lote_id: str
    incluir_avisos: bool = True


async def _procesar(trabajo_id: str, filas: list[dict], payload: dict):
    t = TRABAJOS[trabajo_id]
    pendientes = []   # (correo, nombre, contraseña temporal, user_id) -- solo en memoria
    async with AsyncSessionFactory() as db:
        for f in filas:
            r = {"fila": f["fila"], "nombre": f["nombre"], "correo": f["correo"], "empresa": f["empresa"],
                 "accesos": ", ".join(f"{a['modulo']}: {a['rol']}" for a in f["accesos"]) or "Sin accesos"}
            try:
                u = await user_service.create_user(db, f["company_id"], f["correo"], f["nombre"], matricula=f["matricula"], puesto=f["puesto"],
                                                   departamento=f["departamento"], phone=f["telefono"], requested_by=payload, enviar_bienvenida=False)
                for a in f["accesos"]:
                    await user_service.assign_module_access(db, u["user_id"], a["modulo_id"], a["role_id"], requested_by=payload)
                pendientes.append((f["correo"], f["nombre"], u["temp_password"], u["user_id"]))
                r.update(resultado="creado", detalle="", correo_bienvenida="pendiente")
                t["creados"] += 1
            except Exception as e:
                await db.rollback()
                r.update(resultado="error", detalle=str(getattr(e, "detail", None) or getattr(e, "message", None) or e), correo_bienvenida="no aplica")
                t["errores"] += 1
                log.warning("Alta masiva %s, renglón %s: %s", trabajo_id, f["fila"], r["detalle"])
            t["filas"].append(r)
    t["estado"] = "enviando_correos"
    por_correo = {r["correo"]: r for r in t["filas"]}
    for correo, nombre, clave, uid in pendientes:
        try:
            await user_service._send_welcome_email(correo, nombre, clave, uid)
            por_correo[correo]["correo_bienvenida"] = "enviado"
            t["correos_enviados"] += 1
        except Exception as e:
            log.warning("Alta masiva %s: correo a %s falló: %s", trabajo_id, correo, e)
        await asyncio.sleep(PAUSA_CORREOS)
    pendientes.clear()
    t["estado"] = "terminado"
    t["terminado"] = datetime.now(timezone.utc).isoformat()


@router.post("/confirmar")
async def confirmar(body: Confirmacion, payload=Depends(SOLO_ADMIN)):
    lote = LOTES.pop(body.lote_id, None)
    if not lote:
        raise HTTPException(status_code=404, detail="La revisión expiró o ya se confirmó. Vuelve a subir el archivo.")
    filas = [f for f in lote["filas"] if f["estado"] == "listo" or (body.incluir_avisos and f["estado"] == "aviso")]
    if not filas:
        raise HTTPException(status_code=422, detail="No hay renglones válidos para crear")
    trabajo_id = uuid.uuid4().hex
    TRABAJOS[trabajo_id] = {"estado": "creando", "total": len(filas), "creados": 0, "errores": 0, "correos_enviados": 0,
                            "filas": [], "inicio": datetime.now(timezone.utc).isoformat(), "terminado": None,
                            "segundos_por_correo": PAUSA_CORREOS}
    asyncio.create_task(_procesar(trabajo_id, filas, payload))
    return {"trabajo_id": trabajo_id, "total": len(filas)}


@router.get("/estado/{trabajo_id}")
async def estado(trabajo_id: str, payload=Depends(SOLO_ADMIN)):
    t = TRABAJOS.get(trabajo_id)
    if not t:
        raise HTTPException(status_code=404, detail="No encontré ese proceso (puede que el servicio se haya reiniciado)")
    return t


@router.get("/reporte/{trabajo_id}")
async def reporte(trabajo_id: str, payload=Depends(SOLO_ADMIN)):
    t = TRABAJOS.get(trabajo_id)
    if not t:
        raise HTTPException(status_code=404, detail="No encontré ese proceso")
    wb = Workbook(); ws = wb.active; ws.title = "Resultado"
    titulos = [("Renglón", 10), ("Nombre", 36), ("Correo", 36), ("Empresa", 22), ("Resultado", 12), ("Accesos", 40), ("Correo de bienvenida", 20), ("Detalle", 50)]
    for i, (tt, w) in enumerate(titulos, 1):
        _encabezado(ws, 1, i, tt, True, w)
    for n, r in enumerate(t["filas"], 2):
        for i, v in enumerate([r["fila"], r["nombre"], r["correo"], r["empresa"], r["resultado"], r["accesos"], r["correo_bienvenida"], r["detalle"]], 1):
            c = ws.cell(row=n, column=i, value=v); c.font, c.border = _fuente, _borde
    ws.freeze_panes = "A2"
    return _xlsx(wb, f"Resultado_alta_masiva_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx")
