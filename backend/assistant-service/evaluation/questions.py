# ----------------------------------------------------------------------
# Lectura del conjunto de preguntas de evaluacion
# Cada pregunta trae su verdad de referencia (ground truth): documento y
# ubicacion donde esta la respuesta, anotados por un experto del tema.
# ----------------------------------------------------------------------

# ----------------------------------------------------------------------
# Importaciones
# ----------------------------------------------------------------------
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


# ----------------------------------------------------------------------
# Pregunta con su verdad de referencia
# answerable es falso para las preguntas sin respuesta en el corpus.
# ----------------------------------------------------------------------
@dataclass(frozen=True, slots=True)
class Question:
    id: str
    text: str
    kind: str
    topic: str
    document: str
    location: str

    @property
    def answerable(self) -> bool:
        return self.kind != "sin_respuesta" and self.document not in ("", "ninguno")


# ----------------------------------------------------------------------
# Hoja "Preguntas": filas cuyo ID empieza con P (la fila EJEMPLO no cuenta)
# ----------------------------------------------------------------------
def load_questions(path: Path) -> list[Question]:
    sheet = load_workbook(path, data_only=True)["Preguntas"]
    questions = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        identifier = str(row[0] or "").strip()
        text = str(row[1] or "").strip()
        if not identifier.startswith("P") or not text:
            continue
        clean = [str(value or "").strip() for value in row[2:6]]
        questions.append(Question(identifier, text, *clean))
    return questions
