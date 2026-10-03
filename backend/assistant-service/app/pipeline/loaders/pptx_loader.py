# ----------------------------------------------------------------------
# Cargador de PowerPoint
# Implementa el puerto DocumentLoader para presentaciones .pptx. Entrega
# un bloque por diapositiva (titulo, texto en orden de lectura y notas
# del orador) y un bloque por cada tabla.
# ----------------------------------------------------------------------

# ----------------------------------------------------------------------
# Importaciones
# python-pptx lee la presentacion; MSO_SHAPE_TYPE identifica los grupos
# de formas, que se abren para leer lo que contienen.
# ----------------------------------------------------------------------
from collections.abc import Iterator
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.domain.models import BlockKind, ExtractedBlock, Location, SourceKind
from app.pipeline.loaders.common import normalize_line, table_to_markdown
from app.pipeline.loaders.pptx_rules import ShapeText, build_slide_text, choose_title, reading_order


# ----------------------------------------------------------------------
# Formas de una diapositiva, abriendo los grupos
# ----------------------------------------------------------------------
def _shapes(shapes) -> Iterator:
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _shapes(shape.shapes)
        else:
            yield shape


# ----------------------------------------------------------------------
# Texto de un cuadro: un renglon por parrafo, sin renglones vacios
# ----------------------------------------------------------------------
def _frame_text(shape) -> str:
    lines = (normalize_line(paragraph.text) for paragraph in shape.text_frame.paragraphs)
    return "\n".join(line for line in lines if line)


# ----------------------------------------------------------------------
# Filas de una tabla; las celdas cubiertas por una combinacion van como
# None para que table_to_markdown las una
# ----------------------------------------------------------------------
def _table_rows(shape) -> list[list[str | None]]:
    return [[None if cell.is_spanned else cell.text for cell in row.cells]
            for row in shape.table.rows]


# ----------------------------------------------------------------------
# Notas del orador de una diapositiva, si las tiene
# ----------------------------------------------------------------------
def _notes(slide) -> str:
    if not slide.has_notes_slide:
        return ""
    frame = slide.notes_slide.notes_text_frame
    return frame.text if frame is not None else ""


class PowerPointLoader:
    kind = SourceKind.POWERPOINT

    def supports(self, path: Path) -> bool:
        return path.suffix.lower() == ".pptx"

    # ------------------------------------------------------------------
    # Cada diapositiva: el espacio de titulo se separa del cuerpo; los
    # demas cuadros se ordenan y las tablas se guardan aparte
    # ------------------------------------------------------------------
    def load(self, path: Path) -> list[ExtractedBlock]:
        blocks: list[ExtractedBlock] = []
        for number, slide in enumerate(Presentation(str(path)).slides, start=1):
            title_shape = slide.shapes.title
            title_id = title_shape.shape_id if title_shape is not None else None
            placeholder = title_shape.text_frame.text if title_shape is not None and title_shape.has_text_frame else None

            texts: list[ShapeText] = []
            tables: list[list[list[str | None]]] = []
            for shape in _shapes(slide.shapes):
                if title_id is not None and shape.shape_id == title_id:
                    continue
                if shape.has_table:
                    tables.append(_table_rows(shape))
                elif shape.has_text_frame:
                    text = _frame_text(shape)
                    if text:
                        texts.append(ShapeText(shape.top or 0, shape.left or 0, text))

            ordered = [item.text for item in reading_order(texts)]
            title, body = choose_title(placeholder, ordered)
            path_titles = (title,) if title else ()
            location = Location(slide=number)

            text = build_slide_text(body, _notes(slide))
            if text:
                blocks.append(ExtractedBlock(text=text, location=location,
                                             kind=BlockKind.TEXT, heading_path=path_titles))
            for rows in tables:
                markdown = table_to_markdown(rows)
                if markdown:
                    blocks.append(ExtractedBlock(text=markdown, location=location,
                                                 kind=BlockKind.TABLE, heading_path=path_titles))
        return blocks
